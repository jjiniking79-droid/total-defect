# -*- coding: utf-8 -*-
"""
현재 화면(테이블)에서 '보이는(선택된) 컬럼'만 엑셀로 내보내는 모듈.
"""
import os
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, PatternFill


def export_to_excel(headers, rows, out_path, image_col_index=None, embed_thumbnails=True):
    """
    headers: 내보낼 컬럼 제목 리스트 (보이는 컬럼만)
    rows: 각 행의 값 리스트 (headers와 같은 순서/개수) -> List[List[str]]
          단, image_col_index가 있으면 해당 위치의 값은 이미지 파일 경로(str)여야 함
    out_path: 저장 경로(.xlsx)
    image_col_index: headers 리스트 내에서 썸네일을 넣을 컬럼의 인덱스 (없으면 None)
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "판정결과"

    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = header_fill

    row_height = 60

    for r, row_vals in enumerate(rows, start=2):
        for c, val in enumerate(row_vals, start=1):
            col_idx0 = c - 1
            if image_col_index is not None and col_idx0 == image_col_index and embed_thumbnails:
                img_path = val
                if img_path and os.path.exists(img_path):
                    try:
                        xl_img = XLImage(img_path)
                        xl_img.width = 70
                        xl_img.height = 70
                        cell_ref = f"{get_column_letter(c)}{r}"
                        ws.add_image(xl_img, cell_ref)
                    except Exception:
                        ws.cell(row=r, column=c, value=os.path.basename(str(img_path)))
                else:
                    ws.cell(row=r, column=c, value="")
            else:
                ws.cell(row=r, column=c, value=val)
        ws.row_dimensions[r].height = row_height

    # 열 너비 자동 조정(간단 버전)
    for c, h in enumerate(headers, start=1):
        col_letter = get_column_letter(c)
        if image_col_index is not None and (c - 1) == image_col_index:
            ws.column_dimensions[col_letter].width = 12
        else:
            ws.column_dimensions[col_letter].width = max(12, len(str(h)) + 4)

    ws.freeze_panes = "A2"
    wb.save(out_path)
    return out_path
